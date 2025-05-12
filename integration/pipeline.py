# Complete OpenAI Assistants + Pinecone Implementation

## Folder Structure

```
project_root/
├── knowledge_base/                    # Replace entire folder with new implementation
│   ├── __init__.py
│   ├── config.py
│   ├── openai_assistant_manager.py
│   ├── pinecone_manager.py
│   ├── document_processor.py
│   ├── conversation_manager.py
│   ├── query_engine.py
│   ├── utils/
│   │   ├── __init__.py
│   │   ├── cache_utils.py
│   │   ├── token_counter.py
│   │   └── rate_limiter.py
│   └── exceptions.py
├── integration/
│   ├── kb_integration.py             # Update existing
│   └── pipeline.py                   # Update existing
├── text_to_speech/                   # Keep existing
├── speech_to_text/                   # Keep existing
├── telephony/                        # Keep existing
├── requirements.txt                  # Add new dependencies
└── voice_ai_agent.py                # Update existing
```

## Updated Dependencies (requirements.txt)

Add these new dependencies to your existing requirements.txt:

```python
# Add to existing requirements.txt
openai>=1.3.0
pinecone-client>=2.2.4
tiktoken>=0.5.1
redis>=4.6.0
```

## 1. knowledge_base/__init__.py

```python
"""
OpenAI Assistants API + Pinecone integration for Voice AI Agent.
"""
from knowledge_base.openai_assistant_manager import OpenAIAssistantManager
from knowledge_base.pinecone_manager import PineconeManager
from knowledge_base.document_processor import DocumentProcessor
from knowledge_base.conversation_manager import ConversationManager
from knowledge_base.query_engine import QueryEngine

__version__ = "2.0.0"

__all__ = [
    "OpenAIAssistantManager",
    "PineconeManager", 
    "DocumentProcessor",
    "ConversationManager",
    "QueryEngine",
]
```

## 2. knowledge_base/config.py

```python
"""
Configuration settings for OpenAI + Pinecone knowledge base.
"""
import os
from typing import Dict, Any, List, Optional
from dotenv import load_dotenv

load_dotenv()

# OpenAI Configuration
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
OPENAI_MODEL = os.getenv("OPENAI_MODEL", "gpt-4o")
OPENAI_EMBEDDING_MODEL = os.getenv("OPENAI_EMBEDDING_MODEL", "text-embedding-3-small")
OPENAI_MAX_TOKENS = int(os.getenv("OPENAI_MAX_TOKENS", "4096"))
OPENAI_TEMPERATURE = float(os.getenv("OPENAI_TEMPERATURE", "0.7"))

# Pinecone Configuration
PINECONE_API_KEY = os.getenv("PINECONE_API_KEY")
PINECONE_INDEX_NAME = os.getenv("PINECONE_INDEX_NAME", "voice-ai-knowledge")
PINECONE_ENVIRONMENT = os.getenv("PINECONE_ENVIRONMENT", "us-east-1-aws")
PINECONE_DIMENSION = 1536  # For text-embedding-3-small

# Document Processing
CHUNK_SIZE = int(os.getenv("CHUNK_SIZE", "1000"))
CHUNK_OVERLAP = int(os.getenv("CHUNK_OVERLAP", "200"))
MAX_DOCUMENT_SIZE_MB = int(os.getenv("MAX_DOCUMENT_SIZE_MB", "10"))

# Retrieval Settings
DEFAULT_TOP_K = int(os.getenv("DEFAULT_TOP_K", "5"))
MINIMUM_SIMILARITY_SCORE = float(os.getenv("MINIMUM_SIMILARITY_SCORE", "0.75"))

# Cache Settings
REDIS_URL = os.getenv("REDIS_URL", "redis://localhost:6379/0")
CACHE_TTL = int(os.getenv("CACHE_TTL", "3600"))  # 1 hour

# Rate Limiting
MAX_TOKENS_PER_DAY = int(os.getenv("MAX_TOKENS_PER_DAY", "1000000"))
MAX_REQUESTS_PER_MINUTE = int(os.getenv("MAX_REQUESTS_PER_MINUTE", "60"))

# Supported file types
SUPPORTED_DOCUMENT_TYPES = [
    ".txt", ".md", ".pdf", ".docx", ".doc", 
    ".csv", ".json", ".html", ".htm"
]

def get_openai_config() -> Dict[str, Any]:
    """Get OpenAI configuration."""
    return {
        "api_key": OPENAI_API_KEY,
        "model": OPENAI_MODEL,
        "embedding_model": OPENAI_EMBEDDING_MODEL,
        "max_tokens": OPENAI_MAX_TOKENS,
        "temperature": OPENAI_TEMPERATURE
    }

def get_pinecone_config() -> Dict[str, Any]:
    """Get Pinecone configuration."""
    return {
        "api_key": PINECONE_API_KEY,
        "index_name": PINECONE_INDEX_NAME,
        "environment": PINECONE_ENVIRONMENT,
        "dimension": PINECONE_DIMENSION
    }

def get_processing_config() -> Dict[str, Any]:
    """Get document processing configuration."""
    return {
        "chunk_size": CHUNK_SIZE,
        "chunk_overlap": CHUNK_OVERLAP,
        "max_document_size_mb": MAX_DOCUMENT_SIZE_MB,
        "supported_types": SUPPORTED_DOCUMENT_TYPES
    }

def get_retrieval_config() -> Dict[str, Any]:
    """Get retrieval configuration."""
    return {
        "top_k": DEFAULT_TOP_K,
        "min_similarity": MINIMUM_SIMILARITY_SCORE
    }
```

## 3. knowledge_base/exceptions.py

```python
"""
Custom exceptions for the knowledge base module.
"""

class KnowledgeBaseError(Exception):
    """Base exception for knowledge base errors."""
    pass

class OpenAIError(KnowledgeBaseError):
    """OpenAI API related errors."""
    pass

class PineconeError(KnowledgeBaseError):
    """Pinecone related errors."""
    pass

class DocumentProcessingError(KnowledgeBaseError):
    """Document processing errors."""
    pass

class ConversationError(KnowledgeBaseError):
    """Conversation management errors."""
    pass

class RateLimitError(KnowledgeBaseError):
    """Rate limiting errors."""
    pass

class TokenLimitError(KnowledgeBaseError):
    """Token limit exceeded errors."""
    pass
```

## 4. knowledge_base/utils/token_counter.py

```python
"""
Token counting utilities for OpenAI API.
"""
import tiktoken
import logging
from typing import Dict, Any

logger = logging.getLogger(__name__)

class TokenCounter:
    """Count tokens for OpenAI API calls."""
    
    def __init__(self, model: str = "gpt-4o"):
        """Initialize token counter."""
        self.model = model
        try:
            self.encoding = tiktoken.encoding_for_model(model)
        except KeyError:
            # Fallback to cl100k_base for unknown models
            self.encoding = tiktoken.get_encoding("cl100k_base")
    
    def count_tokens(self, text: str) -> int:
        """Count tokens in text."""
        return len(self.encoding.encode(text))
    
    def count_message_tokens(self, messages: list) -> int:
        """Count tokens in a list of messages."""
        num_tokens = 0
        for message in messages:
            # Every message follows <im_start>{role/name}\n{content}<im_end>\n
            num_tokens += 4
            
            if "role" in message:
                num_tokens += self.count_tokens(message["role"])
                
            if "content" in message:
                num_tokens += self.count_tokens(message["content"])
                
            if "name" in message:
                num_tokens += self.count_tokens(message["name"])
                num_tokens -= 1  # role is omitted when name is present
        
        num_tokens += 2  # every reply is primed with <im_start>assistant
        return num_tokens
    
    def estimate_cost(self, input_tokens: int, output_tokens: int) -> float:
        """Estimate cost based on token usage."""
        # GPT-4o pricing (as of 2024)
        input_cost_per_token = 0.000005  # $5/1M tokens
        output_cost_per_token = 0.000015  # $15/1M tokens
        
        return (input_tokens * input_cost_per_token) + (output_tokens * output_cost_per_token)
```

## 5. knowledge_base/utils/rate_limiter.py

```python
"""
Rate limiting utilities for OpenAI API.
"""
import asyncio
import time
from typing import Dict, Optional
import redis.asyncio as redis
from datetime import datetime, timedelta
import logging

logger = logging.getLogger(__name__)

class RateLimiter:
    """Rate limiting for OpenAI API calls."""
    
    def __init__(self, redis_client: Optional[redis.Redis] = None):
        """Initialize rate limiter."""
        self.redis = redis_client or redis.from_url("redis://localhost:6379/0")
        self.max_requests_per_minute = 60
        self.max_tokens_per_day = 1000000
    
    async def check_rate_limit(self, user_id: str, tokens: int) -> bool:
        """Check if request is within rate limits."""
        current_time = datetime.now()
        
        # Check requests per minute
        minute_key = f"rate_limit:requests:{user_id}:{current_time.strftime('%Y%m%d%H%M')}"
        requests_count = await self.redis.incr(minute_key)
        await self.redis.expire(minute_key, 60)
        
        if requests_count > self.max_requests_per_minute:
            logger.warning(f"Rate limit exceeded for user {user_id}: {requests_count} requests/minute")
            return False
        
        # Check tokens per day
        day_key = f"rate_limit:tokens:{user_id}:{current_time.strftime('%Y%m%d')}"
        token_count = await self.redis.incrby(day_key, tokens)
        await self.redis.expire(day_key, 86400)  # 24 hours
        
        if token_count > self.max_tokens_per_day:
            logger.warning(f"Token limit exceeded for user {user_id}: {token_count} tokens/day")
            return False
        
        return True
    
    async def get_remaining_quota(self, user_id: str) -> Dict[str, int]:
        """Get remaining quota for user."""
        current_time = datetime.now()
        
        # Get current usage
        minute_key = f"rate_limit:requests:{user_id}:{current_time.strftime('%Y%m%d%H%M')}"
        day_key = f"rate_limit:tokens:{user_id}:{current_time.strftime('%Y%m%d')}"
        
        requests_used = int(await self.redis.get(minute_key) or 0)
        tokens_used = int(await self.redis.get(day_key) or 0)
        
        return {
            "requests_remaining": max(0, self.max_requests_per_minute - requests_used),
            "tokens_remaining": max(0, self.max_tokens_per_day - tokens_used)
        }
```

## 6. knowledge_base/utils/cache_utils.py

```python
"""
Caching utilities for knowledge base queries.
"""
import json
import hashlib
import asyncio
from typing import Dict, Any, Optional
import redis.asyncio as redis
import logging

logger = logging.getLogger(__name__)

class CacheManager:
    """Manage caching for knowledge base queries."""
    
    def __init__(self, redis_client: Optional[redis.Redis] = None, ttl: int = 3600):
        """Initialize cache manager."""
        self.redis = redis_client or redis.from_url("redis://localhost:6379/0")
        self.ttl = ttl
    
    def _generate_cache_key(self, query: str, context: Optional[Dict] = None) -> str:
        """Generate cache key for query."""
        # Create hash from query and context
        content = query
        if context:
            content += json.dumps(context, sort_keys=True)
        
        return f"cache:query:{hashlib.md5(content.encode()).hexdigest()}"
    
    async def get(self, query: str, context: Optional[Dict] = None) -> Optional[str]:
        """Get cached response."""
        cache_key = self._generate_cache_key(query, context)
        cached_value = await self.redis.get(cache_key)
        
        if cached_value:
            logger.debug(f"Cache hit for query: {query[:50]}...")
            return json.loads(cached_value)
        
        return None
    
    async def set(self, query: str, response: str, context: Optional[Dict] = None):
        """Cache response."""
        cache_key = self._generate_cache_key(query, context)
        await self.redis.setex(
            cache_key, 
            self.ttl, 
            json.dumps(response)
        )
        logger.debug(f"Cached response for query: {query[:50]}...")
    
    async def invalidate_pattern(self, pattern: str):
        """Invalidate all cache entries matching pattern."""
        keys = await self.redis.keys(f"cache:query:{pattern}*")
        if keys:
            await self.redis.delete(*keys)
            logger.info(f"Invalidated {len(keys)} cache entries")
```

## 7. knowledge_base/openai_assistant_manager.py

```python
"""
OpenAI Assistant management for the knowledge base.
"""
import logging
import asyncio
from typing import Dict, Any, List, Optional, AsyncIterator
from openai import AsyncOpenAI
import json

from knowledge_base.config import get_openai_config
from knowledge_base.exceptions import OpenAIError
from knowledge_base.utils.token_counter import TokenCounter

logger = logging.getLogger(__name__)

class OpenAIAssistantManager:
    """Manage OpenAI Assistants for the voice AI agent."""
    
    def __init__(self):
        """Initialize OpenAI Assistant Manager."""
        self.config = get_openai_config()
        self.client = AsyncOpenAI(api_key=self.config["api_key"])
        self.assistant_id = None
        self.token_counter = TokenCounter(self.config["model"])
        self.system_instructions = self._get_system_instructions()
    
    def _get_system_instructions(self) -> str:
        """Get system instructions for the assistant."""
        return """You are a helpful assistant for a voice AI agent that handles customer inquiries.

IMPORTANT GUIDELINES:
1. Keep responses conversational and concise for voice interactions
2. When you need specific information, use the search_knowledge_base function
3. Always maintain context within the conversation
4. For complex queries, break down information into digestible parts
5. If you don't know something, admit it and offer to search for information
6. Prioritize accuracy and cite sources when using retrieved information

VOICE INTERACTION CONSIDERATIONS:
- Speak naturally, as if in a phone conversation
- Use brief pauses (commas) for better speech synthesis
- Avoid long lists or complex formatting
- Summarize key points clearly"""
    
    async def create_assistant(self) -> str:
        """Create a new OpenAI assistant with function calling capabilities."""
        try:
            assistant = await self.client.beta.assistants.create(
                name="Voice AI Knowledge Assistant",
                instructions=self.system_instructions,
                model=self.config["model"],
                tools=[
                    {
                        "type": "function",
                        "function": {
                            "name": "search_knowledge_base",
                            "description": "Search the knowledge base for relevant information",
                            "parameters": {
                                "type": "object",
                                "properties": {
                                    "query": {
                                        "type": "string",
                                        "description": "The search query"
                                    },
                                    "filters": {
                                        "type": "object",
                                        "description": "Optional metadata filters for the search"
                                    },
                                    "top_k": {
                                        "type": "integer",
                                        "description": "Number of results to return (default: 5)"
                                    }
                                },
                                "required": ["query"]
                            }
                        }
                    }
                ]
            )
            
            self.assistant_id = assistant.id
            logger.info(f"Created OpenAI Assistant: {self.assistant_id}")
            return self.assistant_id
            
        except Exception as e:
            logger.error(f"Error creating assistant: {e}")
            raise OpenAIError(f"Failed to create assistant: {str(e)}")
    
    async def get_or_create_assistant(self) -> str:
        """Get existing assistant or create new one."""
        if not self.assistant_id:
            # Try to find existing assistant first
            assistants = await self.client.beta.assistants.list()
            for assistant in assistants.data:
                if assistant.name == "Voice AI Knowledge Assistant":
                    self.assistant_id = assistant.id
                    logger.info(f"Found existing assistant: {self.assistant_id}")
                    return self.assistant_id
            
            # Create new if not found
            return await self.create_assistant()
        
        return self.assistant_id
    
    async def create_thread(self) -> str:
        """Create a new conversation thread."""
        try:
            thread = await self.client.beta.threads.create()
            logger.debug(f"Created thread: {thread.id}")
            return thread.id
        except Exception as e:
            logger.error(f"Error creating thread: {e}")
            raise OpenAIError(f"Failed to create thread: {str(e)}")
    
    async def add_message_to_thread(self, thread_id: str, message: str, role: str = "user"):
        """Add a message to a thread."""
        try:
            await self.client.beta.threads.messages.create(
                thread_id=thread_id,
                role=role,
                content=message
            )
            logger.debug(f"Added message to thread {thread_id}")
        except Exception as e:
            logger.error(f"Error adding message to thread: {e}")
            raise OpenAIError(f"Failed to add message: {str(e)}")
    
    async def run_assistant(self, thread_id: str, assistant_id: Optional[str] = None) -> AsyncIterator[Dict[str, Any]]:
        """Run the assistant on a thread with streaming support."""
        if not assistant_id:
            assistant_id = await self.get_or_create_assistant()
        
        try:
            # Create run with streaming
            run = await self.client.beta.threads.runs.create(
                thread_id=thread_id,
                assistant_id=assistant_id,
                stream=True
            )
            
            # Process streaming events
            async for event in run:
                if event.event == "thread.message.delta":
                    # Extract text delta
                    if event.data.delta.content:
                        for content in event.data.delta.content:
                            if content.type == "text":
                                yield {
                                    "type": "text_delta",
                                    "content": content.text.value,
                                    "done": False
                                }
                
                elif event.event == "thread.run.requires_action":
                    # Handle function calls
                    tool_calls = event.data.required_action.submit_tool_outputs.tool_calls
                    yield {
                        "type": "function_calls",
                        "tool_calls": tool_calls,
                        "done": False
                    }
                
                elif event.event == "thread.run.completed":
                    yield {
                        "type": "completed",
                        "done": True
                    }
                
                elif event.event == "thread.run.failed":
                    yield {
                        "type": "error",
                        "error": event.data.last_error.message,
                        "done": True
                    }
                    
        except Exception as e:
            logger.error(f"Error running assistant: {e}")
            yield {
                "type": "error",
                "error": str(e),
                "done": True
            }
    
    async def submit_tool_outputs(self, thread_id: str, run_id: str, tool_outputs: List[Dict[str, Any]]):
        """Submit tool outputs for function calls."""
        try:
            await self.client.beta.threads.runs.submit_tool_outputs(
                thread_id=thread_id,
                run_id=run_id,
                tool_outputs=tool_outputs
            )
            logger.debug(f"Submitted tool outputs for run {run_id}")
        except Exception as e:
            logger.error(f"Error submitting tool outputs: {e}")
            raise OpenAIError(f"Failed to submit tool outputs: {str(e)}")
    
    async def get_thread_messages(self, thread_id: str, limit: int = 20) -> List[Dict[str, Any]]:
        """Get messages from a thread."""
        try:
            messages = await self.client.beta.threads.messages.list(
                thread_id=thread_id,
                limit=limit
            )
            
            return [
                {
                    "id": msg.id,
                    "role": msg.role,
                    "content": msg.content[0].text.value if msg.content else "",
                    "created_at": msg.created_at
                }
                for msg in messages.data
            ]
        except Exception as e:
            logger.error(f"Error getting thread messages: {e}")
            raise OpenAIError(f"Failed to get thread messages: {str(e)}")
    
    async def delete_thread(self, thread_id: str):
        """Delete a thread."""
        try:
            await self.client.beta.threads.delete(thread_id)
            logger.debug(f"Deleted thread: {thread_id}")
        except Exception as e:
            logger.error(f"Error deleting thread: {e}")
            # Don't raise error as thread cleanup is not critical
```

## 8. knowledge_base/pinecone_manager.py

```python
"""
Pinecone vector database management.
"""
import logging
import asyncio
from typing import Dict, Any, List, Optional, Tuple
import numpy as np
from pinecone import Pinecone, PodSpec, ServerlessSpec
import uuid

from knowledge_base.config import get_pinecone_config, get_openai_config
from knowledge_base.exceptions import PineconeError
from openai import AsyncOpenAI

logger = logging.getLogger(__name__)

class PineconeManager:
    """Manage Pinecone vector database operations."""
    
    def __init__(self):
        """Initialize Pinecone Manager."""
        self.config = get_pinecone_config()
        self.openai_config = get_openai_config()
        self.client = Pinecone(api_key=self.config["api_key"])
        self.openai_client = AsyncOpenAI(api_key=self.openai_config["api_key"])
        self.index = None
        self.index_name = self.config["index_name"]
        
    async def init(self):
        """Initialize Pinecone index."""
        try:
            # Check if index exists
            existing_indexes = self.client.list_indexes()
            index_names = [idx.name for idx in existing_indexes]
            
            if self.index_name not in index_names:
                # Create index if it doesn't exist
                logger.info(f"Creating Pinecone index: {self.index_name}")
                self.client.create_index(
                    name=self.index_name,
                    dimension=self.config["dimension"],
                    metric="cosine",
                    spec=ServerlessSpec(
                        cloud='aws',
                        region=self.config.get("environment", "us-east-1")
                    )
                )
                
                # Wait for index to be ready
                await asyncio.sleep(5)
            
            # Connect to index
            self.index = self.client.Index(self.index_name)
            logger.info(f"Connected to Pinecone index: {self.index_name}")
            
        except Exception as e:
            logger.error(f"Error initializing Pinecone: {e}")
            raise PineconeError(f"Failed to initialize Pinecone: {str(e)}")
    
    async def generate_embeddings(self, texts: List[str]) -> List[List[float]]:
        """Generate embeddings using OpenAI."""
        try:
            # Process in batches to respect API limits
            batch_size = 50
            all_embeddings = []
            
            for i in range(0, len(texts), batch_size):
                batch = texts[i:i + batch_size]
                
                response = await self.openai_client.embeddings.create(
                    input=batch,
                    model=self.openai_config["embedding_model"]
                )
                
                batch_embeddings = [data.embedding for data in response.data]
                all_embeddings.extend(batch_embeddings)
                
                # Add small delay to respect rate limits
                if i + batch_size < len(texts):
                    await asyncio.sleep(0.1)
            
            return all_embeddings
            
        except Exception as e:
            logger.error(f"Error generating embeddings: {e}")
            raise PineconeError(f"Failed to generate embeddings: {str(e)}")
    
    async def upsert_documents(self, documents: List[Dict[str, Any]]) -> int:
        """Upsert documents to Pinecone."""
        if not self.index:
            await self.init()
        
        try:
            # Extract texts for embedding
            texts = [doc["text"] for doc in documents]
            embeddings = await self.generate_embeddings(texts)
            
            # Prepare vectors for upsert
            vectors = []
            for i, (doc, embedding) in enumerate(zip(documents, embeddings)):
                vector_id = doc.get("id", str(uuid.uuid4()))
                
                vectors.append({
                    "id": vector_id,
                    "values": embedding,
                    "metadata": {
                        "text": doc["text"],
                        "source": doc.get("source", "unknown"),
                        "chunk_index": doc.get("chunk_index", 0),
                        "document_id": doc.get("document_id", ""),
                        **doc.get("metadata", {})
                    }
                })
            
            # Upsert in batches
            batch_size = 100
            upserted_count = 0
            
            for i in range(0, len(vectors), batch_size):
                batch = vectors[i:i + batch_size]
                self.index.upsert(vectors=batch)
                upserted_count += len(batch)
                
                # Add delay between batches
                if i + batch_size < len(vectors):
                    await asyncio.sleep(0.5)
            
            logger.info(f"Upserted {upserted_count} documents to Pinecone")
            return upserted_count
            
        except Exception as e:
            logger.error(f"Error upserting documents: {e}")
            raise PineconeError(f"Failed to upsert documents: {str(e)}")
    
    async def query(
        self, 
        query_text: str, 
        top_k: int = 5,
        filter_dict: Optional[Dict[str, Any]] = None,
        include_metadata: bool = True
    ) -> List[Dict[str, Any]]:
        """Query Pinecone for similar documents."""
        if not self.index:
            await self.init()
        
        try:
            # Generate embedding for query
            embeddings = await self.generate_embeddings([query_text])
            query_embedding = embeddings[0]
            
            # Query Pinecone
            results = self.index.query(
                vector=query_embedding,
                top_k=top_k,
                filter=filter_dict,
                include_metadata=include_metadata,
                include_values=False
            )
            
            # Format results
            formatted_results = []
            for match in results.matches:
                result = {
                    "id": match.id,
                    "score": match.score,
                    "metadata": match.metadata
                }
                
                if include_metadata and match.metadata:
                    result["text"] = match.metadata.get("text", "")
                    result["source"] = match.metadata.get("source", "unknown")
                
                formatted_results.append(result)
            
            logger.debug(f"Found {len(formatted_results)} results for query")
            return formatted_results
            
        except Exception as e:
            logger.error(f"Error querying Pinecone: {e}")
            raise PineconeError(f"Failed to query Pinecone: {str(e)}")
    
    async def delete_documents(self, document_ids: List[str]):
        """Delete documents from Pinecone."""
        if not self.index:
            await self.init()
        
        try:
            self.index.delete(ids=document_ids)
            logger.info(f"Deleted {len(document_ids)} documents from Pinecone")
        except Exception as e:
            logger.error(f"Error deleting documents: {e}")
            raise PineconeError(f"Failed to delete documents: {str(e)}")
    
    async def get_stats(self) -> Dict[str, Any]:
        """Get Pinecone index statistics."""
        if not self.index:
            await self.init()
        
        try:
            stats = self.index.describe_index_stats()
            return {
                "total_vectors": stats.total_vector_count,
                "index_fullness": stats.index_fullness,
                "dimension": stats.dimension,
                "namespaces": stats.namespaces
            }
        except Exception as e:
            logger.error(f"Error getting stats: {e}")
            return {}
```

## 9. knowledge_base/document_processor.py

```python
"""
Document processing for the knowledge base.
"""
import logging
import os
from typing import List, Dict, Any, Optional
from pathlib import Path
import hashlib
import tiktoken

from knowledge_base.config import get_processing_config
from knowledge_base.exceptions import DocumentProcessingError

logger = logging.getLogger(__name__)

class DocumentProcessor:
    """Process documents for the knowledge base."""
    
    def __init__(self):
        """Initialize document processor."""
        self.config = get_processing_config()
        self.encoding = tiktoken.get_encoding("cl100k_base")
    
    def process_text(self, text: str, source: str = "unknown") -> List[Dict[str, Any]]:
        """Process text into chunks suitable for vector storage."""
        try:
            # Split text into chunks
            chunks = self._split_text(text)
            
            # Create document objects
            documents = []
            for i, chunk in enumerate(chunks):
                doc_id = self._generate_document_id(chunk, source, i)
                
                documents.append({
                    "id": doc_id,
                    "text": chunk,
                    "source": source,
                    "chunk_index": i,
                    "chunk_count": len(chunks),
                    "metadata": {
                        "token_count": self._count_tokens(chunk),
                        "word_count": len(chunk.split()),
                        "char_count": len(chunk)
                    }
                })
            
            logger.info(f"Processed text into {len(documents)} chunks")
            return documents
            
        except Exception as e:
            logger.error(f"Error processing text: {e}")
            raise DocumentProcessingError(f"Failed to process text: {str(e)}")
    
    def process_file(self, file_path: str) -> List[Dict[str, Any]]:
        """Process a file into document chunks."""
        try:
            # Check file exists and is supported
            if not os.path.exists(file_path):
                raise DocumentProcessingError(f"File not found: {file_path}")
            
            file_ext = Path(file_path).suffix.lower()
            if file_ext not in self.config["supported_types"]:
                raise DocumentProcessingError(f"Unsupported file type: {file_ext}")
            
            # Check file size
            file_size = os.path.getsize(file_path)
            max_size = self.config["max_document_size_mb"] * 1024 * 1024
            if file_size > max_size:
                raise DocumentProcessingError(f"File too large: {file_size} bytes")
            
            # Read file content based on type
            if file_ext == ".txt":
                text = self._read_text_file(file_path)
            elif file_ext == ".pdf":
                text = self._read_pdf_file(file_path)
            elif file_ext == ".docx":
                text = self._read_docx_file(file_path)
            else:
                text = self._read_text_file(file_path)  # Fallback
            
            # Process the extracted text
            source = Path(file_path).name
            return self.process_text(text, source)
            
        except Exception as e:
            logger.error(f"Error processing file {file_path}: {e}")
            raise DocumentProcessingError(f"Failed to process file: {str(e)}")
    
    def _split_text(self, text: str) -> List[str]:
        """Split text into chunks with overlap."""
        chunk_size = self.config["chunk_size"]
        chunk_overlap = self.config["chunk_overlap"]
        
        # Simple token-based chunking
        tokens = self.encoding.encode(text)
        chunks = []
        
        for i in range(0, len(tokens), chunk_size - chunk_overlap):
            chunk_tokens = tokens[i:i + chunk_size]
            chunk_text = self.encoding.decode(chunk_tokens)
            chunks.append(chunk_text)
        
        # Ensure we have at least one chunk
        if not chunks and text:
            chunks = [text]
        
        return chunks
    
    def _count_tokens(self, text: str) -> int:
        """Count tokens in text."""
        return len(self.encoding.encode(text))
    
    def _generate_document_id(self, text: str, source: str, chunk_index: int) -> str:
        """Generate unique document ID."""
        content = f"{source}_{chunk_index}_{text[:100]}"
        return hashlib.md5(content.encode()).hexdigest()
    
    def _read_text_file(self, file_path: str) -> str:
        """Read plain text file."""
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    
    def _read_pdf_file(self, file_path: str) -> str:
        """Read PDF file."""
        try:
            import PyPDF2
            text = ""
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text += page.extract_text() + "\n"
            return text
        except ImportError:
            logger.warning("PyPDF2 not installed, falling back to text reading")
            return self._read_text_file(file_path)
    
    def _read_docx_file(self, file_path: str) -> str:
        """Read DOCX file."""
        try:
            from docx import Document
            doc = Document(file_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text
        except ImportError:
            logger.warning("python-docx not installed, falling back to text reading")
            return self._read_text_file(file_path)
```

## 10. knowledge_base/conversation_manager.py

```python
"""
Conversation management using OpenAI Assistants API.
"""
import logging
import asyncio
from typing import Dict, Any, List, Optional, AsyncIterator
import json

from knowledge_base.openai_assistant_manager import OpenAIAssistantManager
from knowledge_base.pinecone_manager import PineconeManager
from knowledge_base.exceptions import ConversationError
from knowledge_base.utils.cache_utils import CacheManager
from knowledge_base.utils.rate_limiter import RateLimiter

logger = logging.getLogger(__name__)

class ConversationManager:
    """Manage conversations using OpenAI Assistants API."""
    
    def __init__(self):
        """Initialize conversation manager."""
        self.openai_manager = OpenAIAssistantManager()
        self.pinecone_manager = PineconeManager()
        self.cache = CacheManager()
        self.rate_limiter = RateLimiter()
        self.user_threads: Dict[str, str] = {}
    
    async def init(self):
        """Initialize all components."""
        try:
            await self.openai_manager.get_or_create_assistant()
            await self.pinecone_manager.init()
            logger.info("Conversation manager initialized successfully")
        except Exception as e:
            logger.error(f"Error initializing conversation manager: {e}")
            raise ConversationError(f"Failed to initialize: {str(e)}")
    
    async def get_or_create_thread(self, user_id: str) -> str:
        """Get existing thread or create new one for user."""
        if user_id not in self.user_threads:
            thread_id = await self.openai_manager.create_thread()
            self.user_threads[user_id] = thread_id
            logger.debug(f"Created new thread for user {user_id}: {thread_id}")
        
        return self.user_threads[user_id]
    
    async def handle_user_input(self, user_id: str, message: str) -> Dict[str, Any]:
        """Handle user input and return response."""
        try:
            # Check rate limits
            if not await self.rate_limiter.check_rate_limit(user_id, 1000):
                return {
                    "response": "You've reached your rate limit. Please try again later.",
                    "error": "rate_limit_exceeded"
                }
            
            # Check cache first
            cached_response = await self.cache.get(message, {"user_id": user_id})
            if cached_response:
                return {
                    "response": cached_response,
                    "cached": True
                }
            
            # Get or create thread for user
            thread_id = await self.get_or_create_thread(user_id)
            
            # Add user message to thread
            await self.openai_manager.add_message_to_thread(thread_id, message)
            
            # Process the message and get response
            response = await self._process_message(thread_id, user_id)
            
            # Cache the response
            if response and not response.get("error"):
                await self.cache.set(message, response["response"], {"user_id": user_id})
            
            return response
            
        except Exception as e:
            logger.error(f"Error handling user input: {e}")
            return {
                "response": "I'm sorry, I encountered an error processing your request.",
                "error": str(e)
            }
    
    async def handle_user_input_streaming(self, user_id: str, message: str) -> AsyncIterator[Dict[str, Any]]:
        """Handle user input with streaming response."""
        try:
            # Check rate limits
            if not await self.rate_limiter.check_rate_limit(user_id, 1000):
                yield {
                    "chunk": "You've reached your rate limit. Please try again later.",
                    "done": True,
                    "error": "rate_limit_exceeded"
                }
                return
            
            # Get or create thread for user
            thread_id = await self.get_or_create_thread(user_id)
            
            # Add user message to thread
            await self.openai_manager.add_message_to_thread(thread_id, message)
            
            # Process with streaming
            full_response = ""
            pending_function_calls = []
            run_id = None
            
            async for event in self.openai_manager.run_assistant(thread_id):
                if event["type"] == "text_delta":
                    chunk = event["content"]
                    full_response += chunk
                    yield {
                        "chunk": chunk,
                        "done": False
                    }
                
                elif event["type"] == "function_calls":
                    pending_function_calls = event["tool_calls"]
                    run_id = event.get("run_id")
                    
                    # Process function calls
                    tool_outputs = await self._process_function_calls(pending_function_calls)
                    
                    # Submit tool outputs
                    if run_id:
                        await self.openai_manager.submit_tool_outputs(thread_id, run_id, tool_outputs)
                
                elif event["type"] == "completed":
                    yield {
                        "chunk": "",
                        "done": True,
                        "full_response": full_response
                    }
                    break
                
                elif event["type"] == "error":
                    yield {
                        "chunk": f"Error: {event['error']}",
                        "done": True,
                        "error": event["error"]
                    }
                    break
            
            # Cache the full response
            if full_response:
                await self.cache.set(message, full_response, {"user_id": user_id})
                
        except Exception as e:
            logger.error(f"Error handling streaming input: {e}")
            yield {
                "chunk": "I'm sorry, I encountered an error processing your request.",
                "done": True,
                "error": str(e)
            }
    
    async def _process_message(self, thread_id: str, user_id: str) -> Dict[str, Any]:
        """Process message with non-streaming response."""
        try:
            response_text = ""
            run_completed = False
            
            async for event in self.openai_manager.run_assistant(thread_id):
                if event["type"] == "text_delta":
                    response_text += event["content"]
                
                elif event["type"] == "function_calls":
                    # Process function calls
                    tool_outputs = await self._process_function_calls(event["tool_calls"])
                    
                    # Submit tool outputs (this will continue the run)
                    await self.openai_manager.submit_tool_outputs(
                        thread_id, 
                        event.get("run_id"), 
                        tool_outputs
                    )
                
                elif event["type"] == "completed":
                    run_completed = True
                    break
                
                elif event["type"] == "error":
                    return {
                        "response": "I encountered an error processing your request.",
                        "error": event["error"]
                    }
            
            if run_completed and response_text:
                return {
                    "response": response_text,
                    "user_id": user_id
                }
            else:
                return {
                    "response": "I'm sorry, I couldn't generate a response.",
                    "error": "no_response_generated"
                }
                
        except Exception as e:
            logger.error(f"Error processing message: {e}")
            return {
                "response": "I'm sorry, I encountered an error.",
                "error": str(e)
            }
    
    async def _process_function_calls(self, tool_calls: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """Process function calls from the assistant."""
        tool_outputs = []
        
        for tool_call in tool_calls:
            if tool_call.function.name == "search_knowledge_base":
                try:
                    # Parse arguments
                    arguments = json.loads(tool_call.function.arguments)
                    query = arguments.get("query", "")
                    filters = arguments.get("filters")
                    top_k = arguments.get("top_k", 5)
                    
                    # Search Pinecone
                    results = await self.pinecone_manager.query(
                        query_text=query,
                        top_k=top_k,
                        filter_dict=filters
                    )
                    
                    # Format results for the assistant
                    formatted_results = self._format_search_results(results)
                    
                    tool_outputs.append({
                        "tool_call_id": tool_call.id,
                        "output": formatted_results
                    })
                    
                except Exception as e:
                    logger.error(f"Error in search_knowledge_base: {e}")
                    tool_outputs.append({
                        "tool_call_id": tool_call.id,
                        "output": f"Error searching knowledge base: {str(e)}"
                    })
        
        return tool_outputs
    
    def _format_search_results(self, results: List[Dict[str, Any]]) -> str:
        """Format search results for the assistant."""
        if not results:
            return "No relevant information found in the knowledge base."
        
        formatted = "Found the following relevant information:\n\n"
        
        for i, result in enumerate(results, 1):
            formatted += f"**Result {i}** (Relevance: {result['score']:.3f})\n"
            formatted += f"Source: {result.get('source', 'Unknown')}\n"
            formatted += f"Content: {result.get('text', 'No content available')}\n\n"
        
        return formatted
    
    async def get_conversation_history(self, user_id: str, limit: int = 10) -> List[Dict[str, Any]]:
        """Get conversation history for a user."""
        if user_id not in self.user_threads:
            return []
        
        thread_id = self.user_threads[user_id]
        return await self.openai_manager.get_thread_messages(thread_id, limit)
    
    async def reset_conversation(self, user_id: str):
        """Reset conversation for a user."""
        if user_id in self.user_threads:
            # Delete old thread
            await self.openai_manager.delete_thread(self.user_threads[user_id])
            
            # Create new thread
            self.user_threads[user_id] = await self.openai_manager.create_thread()
            
            logger.info(f"Reset conversation for user {user_id}")
```

## 11. knowledge_base/query_engine.py

```python
"""
Query engine combining OpenAI Assistants and Pinecone.
"""
import logging
import asyncio
from typing import Dict, Any, List, Optional, AsyncIterator

from knowledge_base.pinecone_manager import PineconeManager
from knowledge_base.openai_assistant_manager import OpenAIAssistantManager
from knowledge_base.config import get_retrieval_config
from knowledge_base.exceptions import KnowledgeBaseError

logger = logging.getLogger(__name__)

class QueryEngine:
    """Query engine using OpenAI Assistants and Pinecone."""
    
    def __init__(self):
        """Initialize query engine."""
        self.pinecone_manager = PineconeManager()
        self.openai_manager = OpenAIAssistantManager()
        self.config = get_retrieval_config()
    
    async def init(self):
        """Initialize all components."""
        await self.pinecone_manager.init()
        await self.openai_manager.get_or_create_assistant()
        logger.info("Query engine initialized")
    
    async def query(self, query_text: str, user_id: Optional[str] = None) -> Dict[str, Any]:
        """Query the knowledge base."""
        try:
            # Step 1: Search Pinecone for relevant documents
            search_results = await self.pinecone_manager.query(
                query_text=query_text,
                top_k=self.config["top_k"]
            )
            
            # Step 2: Format context from search results
            context = self._format_context(search_results)
            
            # Step 3: Create a temporary thread for this query
            thread_id = await self.openai_manager.create_thread()
            
            # Step 4: Add query with context to thread
            message_with_context = f"""Question: {query_text}

Relevant information from knowledge base:
{context}

Please answer the question based on the provided information. If the information is not sufficient, please say so."""

            await self.openai_manager.add_message_to_thread(thread_id, message_with_context)
            
            # Step 5: Get response from assistant
            response_text = ""
            async for event in self.openai_manager.run_assistant(thread_id):
                if event["type"] == "text_delta":
                    response_text += event["content"]
                elif event["type"] == "completed":
                    break
                elif event["type"] == "error":
                    return {
                        "query": query_text,
                        "response": "Error processing your query.",
                        "error": event["error"]
                    }
            
            # Clean up thread
            await self.openai_manager.delete_thread(thread_id)
            
            return {
                "query": query_text,
                "response": response_text,
                "sources": search_results,
                "user_id": user_id
            }
            
        except Exception as e:
            logger.error(f"Error in query: {e}")
            return {
                "query": query_text,
                "response": "I'm sorry, I encountered an error processing your query.",
                "error": str(e)
            }
    
    async def query_with_streaming(self, query_text: str, user_id: Optional[str] = None) -> AsyncIterator[Dict[str, Any]]:
        """Query with streaming response."""
        try:
            # Step 1: Search Pinecone for relevant documents
            search_results = await self.pinecone_manager.query(
                query_text=query_text,
                top_k=self.config["top_k"]
            )
            
            # Step 2: Format context from search results
            context = self._format_context(search_results)
            
            # Step 3: Create a temporary thread for this query
            thread_id = await self.openai_manager.create_thread()
            
            # Step 4: Add query with context to thread
            message_with_context = f"""Question: {query_text}

Relevant information from knowledge base:
{context}

Please answer the question based on the provided information. If the information is not sufficient, please say so."""

            await self.openai_manager.add_message_to_thread(thread_id, message_with_context)
            
            # Step 5: Stream response from assistant
            full_response = ""
            async for event in self.openai_manager.run_assistant(thread_id):
                if event["type"] == "text_delta":
                    chunk = event["content"]
                    full_response += chunk
                    yield {
                        "chunk": chunk,
                        "done": False,
                        "sources": search_results
                    }
                elif event["type"] == "completed":
                    yield {
                        "chunk": "",
                        "done": True,
                        "full_response": full_response,
                        "sources": search_results
                    }
                    break
                elif event["type"] == "error":
                    yield {
                        "chunk": f"Error: {event['error']}",
                        "done": True,
                        "error": event["error"]
                    }
                    break
            
            # Clean up thread
            await self.openai_manager.delete_thread(thread_id)
            
        except Exception as e:
            logger.error(f"Error in streaming query: {e}")
            yield {
                "chunk": "I'm sorry, I encountered an error processing your query.",
                "done": True,
                "error": str(e)
            }
    
    async def retrieve_with_sources(self, query_text: str, top_k: Optional[int] = None) -> Dict[str, Any]:
        """Retrieve documents with source information."""
        try:
            results = await self.pinecone_manager.query(
                query_text=query_text,
                top_k=top_k or self.config["top_k"]
            )
            
            formatted_results = []
            sources = set()
            
            for result in results:
                formatted_result = {
                    "id": result["id"],
                    "text": result.get("text", ""),
                    "score": result["score"],
                    "metadata": result.get("metadata", {})
                }
                formatted_results.append(formatted_result)
                
                # Extract unique sources
                source = result.get("source", "unknown")
                sources.add(source)
            
            return {
                "query": query_text,
                "results": formatted_results,
                "sources": list(sources)
            }
            
        except Exception as e:
            logger.error(f"Error retrieving documents: {e}")
            return {
                "query": query_text,
                "results": [],
                "sources": [],
                "error": str(e)
            }
    
    def format_retrieved_context(self, results: List[Dict[str, Any]]) -> str:
        """Format retrieved context for use in prompts."""
        return self._format_context(results)
    
    def _format_context(self, results: List[Dict[str, Any]]) -> str:
        """Format search results into context string."""
        if not results:
            return "No relevant information found."
        
        context_parts = []
        for i, result in enumerate(results, 1):
            text = result.get("text", "")
            source = result.get("source", "Unknown")
            score = result.get("score", 0)
            
            context_parts.append(f"""
Document {i} (Source: {source}, Relevance: {score:.3f}):
{text}
""")
        
        return "\n".join(context_parts)
    
    async def get_stats(self) -> Dict[str, Any]:
        """Get query engine statistics."""
        try:
            pinecone_stats = await self.pinecone_manager.get_stats()
            return {
                "pinecone_stats": pinecone_stats,
                "config": self.config
            }
        except Exception as e:
            logger.error(f"Error getting stats: {e}")
            return {"error": str(e)}
```

## Updated integration/kb_integration.py

```python
"""
Knowledge Base integration for Voice AI Agent - Updated for OpenAI + Pinecone.
"""
import logging
import time
import asyncio
from typing import Optional, Dict, Any, AsyncIterator, List

from knowledge_base.conversation_manager import ConversationManager
from knowledge_base.query_engine import QueryEngine
from knowledge_base.utils.cache_utils import CacheManager

logger = logging.getLogger(__name__)

class KnowledgeBaseIntegration:
    """
    Knowledge Base integration using OpenAI Assistants and Pinecone.
    """
    
    def __init__(
        self,
        user_id: Optional[str] = None,
        use_cache: bool = True
    ):
        """
        Initialize the Knowledge Base integration.
        
        Args:
            user_id: User identifier for conversation tracking
            use_cache: Whether to use response caching
        """
        self.user_id = user_id
        self.use_cache = use_cache
        
        self.conversation_manager = ConversationManager()
        self.query_engine = QueryEngine()
        self.initialized = False
        
        if self.use_cache:
            self.response_cache = CacheManager()
    
    async def init(self):
        """Initialize the knowledge base integration."""
        try:
            await self.conversation_manager.init()
            await self.query_engine.init()
            self.initialized = True
            logger.info("Knowledge Base integration initialized successfully")
        except Exception as e:
            logger.error(f"Error initializing Knowledge Base integration: {e}")
            raise
    
    async def query(self, text: str, include_context: bool = False) -> Dict[str, Any]:
        """
        Query the knowledge base and generate a response.
        
        Args:
            text: Query text
            include_context: Whether to include context in the response
            
        Returns:
            Dictionary with query response information
        """
        if not self.initialized:
            logger.error("Knowledge Base integration not properly initialized")
            return {"error": "Knowledge Base integration not initialized"}
        
        # Track timing
        start_time = time.time()
        
        # Check cache first if enabled
        if self.use_cache:
            cached_result = await self.response_cache.get(text, {"user_id": self.user_id})
            if cached_result:
                logger.info(f"Cache hit for query: '{text[:50]}...'")
                return {
                    **cached_result,
                    "total_time": time.time() - start_time,
                    "cache_hit": True
                }
        
        try:
            # Handle conversation through OpenAI Assistant
            result = await self.conversation_manager.handle_user_input(
                user_id=self.user_id or "default_user",
                message=text
            )
            
            # Get retrieval results if context is needed
            context_data = {}
            if include_context:
                retrieval_results = await self.query_engine.retrieve_with_sources(text)
                context_data = {
                    "context": self.query_engine.format_retrieved_context(retrieval_results["results"]),
                    "sources": retrieval_results["sources"]
                }
            
            # Prepare final result
            final_result = {
                "query": text,
                "response": result.get("response", ""),
                "total_time": time.time() - start_time,
                **context_data
            }
            
            # Cache result if enabled and no error
            if self.use_cache and not result.get("error"):
                await self.response_cache.set(text, final_result, {"user_id": self.user_id})
            
            return final_result
            
        except Exception as e:
            logger.error(f"Error in Knowledge Base query: {e}")
            return {
                "error": str(e),
                "query": text,
                "total_time": time.time() - start_time
            }
    
    async def query_streaming(self, text: str) -> AsyncIterator[Dict[str, Any]]:
        """
        Query the knowledge base with streaming response generation.
        
        Args:
            text: Query text
            
        Yields:
            Response chunks
        """
        if not self.initialized:
            logger.error("Knowledge Base integration not properly initialized")
            yield {"error": "Knowledge Base integration not initialized", "done": True}
            return
        
        # Check cache for fast responses
        if self.use_cache:
            cached_result = await self.response_cache.get(text, {"user_id": self.user_id})
            if cached_result and "response" in cached_result:
                # For cached responses, yield the entire response as a single chunk
                yield {
                    "chunk": cached_result["response"],
                    "done": False,
                    "cache_hit": True
                }
                yield {
                    "chunk": "",
                    "full_response": cached_result["response"],
                    "done": True,
                    "cache_hit": True,
                    "sources": cached_result.get("sources", [])
                }
                return
        
        try:
            # Stream response from conversation manager
            full_response = ""
            async for chunk in self.conversation_manager.handle_user_input_streaming(
                user_id=self.user_id or "default_user",
                message=text
            ):
                # Pass through chunks directly
                yield chunk
                
                # Accumulate full response for caching
                if not chunk.get("done", False):
                    full_response += chunk.get("chunk", "")
                elif chunk.get("done", False) and chunk.get("full_response"):
                    full_response = chunk["full_response"]
                    
            # Cache the full response
            if self.use_cache and full_response:
                await self.response_cache.set(text, {
                    "response": full_response,
                    "query": text
                }, {"user_id": self.user_id})
                
        except Exception as e:
            logger.error(f"Error in streaming Knowledge Base query: {e}")
            yield {
                "error": str(e),
                "done": True
            }
    
    def reset_conversation(self) -> None:
        """Reset the conversation state."""
        asyncio.create_task(self.conversation_manager.reset_conversation(
            user_id=self.user_id or "default_user"
        ))
    
    def get_conversation_history(self) -> List[Dict[str, Any]]:
        """
        Get conversation history.
        
        Returns:
            List of conversation turns
        """
        # This would need to be implemented with async context
        # For now, return empty list
        logger.warning("get_conversation_history needs async implementation")
        return []
    
    async def get_stats(self) -> Dict[str, Any]:
        """
        Get statistics about the knowledge base.
        
        Returns:
            Dictionary with statistics
        """
        if not self.initialized:
            return {"error": "Not initialized"}
        
        try:
            # Get stats from query engine
            kb_stats = await self.query_engine.get_stats()
            
            # Add integration-specific stats
            stats = {
                "kb_stats": kb_stats,
                "user_id": self.user_id,
                "cache_enabled": self.use_cache
            }
            
            # Add cache stats if enabled
            if self.use_cache:
                stats["cache_stats"] = {
                    "cache_enabled": True,
                    "cache_type": "Redis"
                }
            
            return stats
            
        except Exception as e:
            logger.error(f"Error getting stats: {e}")
            return {"error": str(e)}
```

## Updated voice_ai_agent.py

```python
"""
Voice AI Agent main class - Updated for OpenAI + Pinecone integration.
"""
import os
import logging
import asyncio
import time
from typing import Optional, Dict, Any, Union, Callable, Awaitable
import numpy as np

# Speech processing imports (unchanged)
from speech_to_text.google_cloud_stt import GoogleCloudStreamingSTT
from speech_to_text.stt_integration import STTIntegration

# New knowledge base imports
from knowledge_base.conversation_manager import ConversationManager
from knowledge_base.query_engine import QueryEngine
from knowledge_base.document_processor import DocumentProcessor
from knowledge_base.pinecone_manager import PineconeManager

# ElevenLabs TTS imports (unchanged)
from text_to_speech import ElevenLabsTTS

logger = logging.getLogger(__name__)

class VoiceAIAgent:
    """Main Voice AI Agent class - Updated for OpenAI Assistants + Pinecone."""
    
    def __init__(
        self,
        storage_dir: str = './storage',
        **kwargs
    ):
        """
        Initialize the Voice AI Agent with OpenAI + Pinecone.
        
        Args:
            storage_dir: Directory for persistent storage
            **kwargs: Additional parameters for customization
        """
        self.storage_dir = storage_dir
        
        # STT Parameters
        self.stt_language = kwargs.get('language', 'en-US')
        self.stt_keywords = kwargs.get('keywords', ['price', 'plan', 'cost', 'subscription', 'service'])
        self.enhanced_model = kwargs.get('enhanced_model', True)
        
        # TTS Parameters for ElevenLabs
        self.elevenlabs_api_key = kwargs.get('elevenlabs_api_key', os.getenv('ELEVENLABS_API_KEY'))
        self.elevenlabs_voice_id = kwargs.get('elevenlabs_voice_id', os.getenv('TTS_VOICE_ID', 'EXAVITQu4vr4xnSDxMaL'))
        self.elevenlabs_model_id = kwargs.get('elevenlabs_model_id', os.getenv('TTS_MODEL_ID', 'eleven_turbo_v2'))
        
        # Component placeholders
        self.speech_recognizer = None
        self.stt_integration = None
        self.conversation_manager = None
        self.query_engine = None
        self.tts_client = None
        self.document_processor = None
        self.pinecone_manager = None
        
        # Audio processing
        self.noise_floor = 0.005
        self.noise_samples = []
        self.max_noise_samples = 20
    
    async def init(self):
        """Initialize all components with OpenAI + Pinecone."""
        logger.info("Initializing Voice AI Agent with OpenAI + Pinecone...")
        
        # Initialize speech recognizer with Google Cloud
        self.speech_recognizer = GoogleCloudStreamingSTT(
            language=self.stt_language,
            sample_rate=16000,
            encoding="LINEAR16",
            channels=1,
            interim_results=True,
            speech_context_phrases=self.stt_keywords,
            enhanced_model=self.enhanced_model
        )
        
        # Initialize STT integration 
        self.stt_integration = STTIntegration(
            speech_recognizer=self.speech_recognizer,
            language=self.stt_language
        )
        
        # Initialize new knowledge base components
        self.conversation_manager = ConversationManager()
        await self.conversation_manager.init()
        
        self.query_engine = QueryEngine()
        await self.query_engine.init()
        
        self.document_processor = DocumentProcessor()
        
        self.pinecone_manager = PineconeManager()
        await self.pinecone_manager.init()
        
        # Initialize ElevenLabs TTS client
        try:
            if not self.elevenlabs_api_key:
                raise ValueError("ElevenLabs API key is required")
                
            self.tts_client = ElevenLabsTTS(
                api_key=self.elevenlabs_api_key,
                voice_id=self.elevenlabs_voice_id,
                model_id=self.elevenlabs_model_id,
                container_format="mulaw",
                sample_rate=8000,
                optimize_streaming_latency=4
            )
            
            logger.info(f"Initialized ElevenLabs TTS with voice ID: {self.elevenlabs_voice_id}")
        except Exception as e:
            logger.error(f"Error initializing ElevenLabs TTS: {e}")
            raise
        
        logger.info("Voice AI Agent initialization complete")
    
    async def process_audio(
        self,
        audio_data: Union[bytes, np.ndarray],
        user_id: Optional[str] = None,
        callback: Optional[Callable[[Any], Awaitable[None]]] = None
    ) -> Dict[str, Any]:
        """
        Process audio data with Google Cloud STT and OpenAI response.
        
        Args:
            audio_data: Audio data as numpy array or bytes
            user_id: User identifier for conversation tracking
            callback: Optional callback function
            
        Returns:
            Processing result
        """
        if not self.initialized:
            raise RuntimeError("Voice AI Agent not initialized")
        
        # Process audio for better speech recognition
        if isinstance(audio_data, np.ndarray):
            audio_data = self._process_audio(audio_data)
        
        # Use STT integration for processing
        result = await self.stt_integration.transcribe_audio_data(audio_data, callback=callback)
        
        # Only process valid transcriptions
        if result.get("is_valid", False) and result.get("transcription"):
            transcription = result["transcription"]
            logger.info(f"Valid transcription: {transcription}")
            
            # Process through conversation manager
            response = await self.conversation_manager.handle_user_input(
                user_id=user_id or "default_user",
                message=transcription
            )
            
            # Generate speech using ElevenLabs TTS
            if response and response.get("response"):
                try:
                    speech_audio = await self.tts_client.synthesize(response["response"])
                    return {
                        "transcription": transcription,
                        "response": response.get("response", ""),
                        "speech_audio": speech_audio,
                        "status": "success"
                    }
                except Exception as e:
                    logger.error(f"Error synthesizing speech: {e}")
                    return {
                        "transcription": transcription,
                        "response": response.get("response", ""),
                        "error": f"Speech synthesis error: {str(e)}",
                        "status": "tts_error"
                    }
            else:
                return {
                    "transcription": transcription,
                    "response": response.get("response", ""),
                    "status": "success"
                }
        else:
            logger.info("Invalid or empty transcription")
            return {
                "status": "invalid_transcription",
                "transcription": result.get("transcription", ""),
                "error": "No valid speech detected"
            }
    
    async def process_streaming_audio(
        self,
        audio_stream,
        user_id: Optional[str] = None,
        result_callback: Optional[Callable[[Dict[str, Any]], Awaitable[None]]] = None
    ) -> Dict[str, Any]:
        """
        Process streaming audio with real-time response.
        
        Args:
            audio_stream: Async iterator of audio chunks
            user_id: User identifier
            result_callback: Callback for streaming results
            
        Returns:
            Final processing stats
        """
        if not self.initialized:
            raise RuntimeError("Voice AI Agent not initialized")
            
        # Track stats
        start_time = time.time()
        chunks_processed = 0
        results_count = 0
        
        # Start streaming session
        await self.speech_recognizer.start_streaming()
        
        try:
            # Process each audio chunk
            async for chunk in audio_stream:
                chunks_processed += 1
                
                # Process audio for better recognition
                if isinstance(chunk, np.ndarray):
                    chunk = self._process_audio(chunk)
                
                # Process through Google Cloud STT
                async def process_result(result):
                    # Only handle final results
                    if result.is_final:
                        transcription = self.stt_integration.cleanup_transcription(result.text)
                        
                        # Process if valid
                        if transcription and self.stt_integration.is_valid_transcription(transcription):
                            # Get response from conversation manager
                            response = await self.conversation_manager.handle_user_input(
                                user_id=user_id or "default_user",
                                message=transcription
                            )
                            
                            # Generate speech with ElevenLabs TTS
                            speech_audio = None
                            tts_error = None
                            
                            if response and response.get("response"):
                                try:
                                    speech_audio = await self.tts_client.synthesize(response["response"])
                                except Exception as e:
                                    logger.error(f"Error synthesizing speech: {e}")
                                    tts_error = str(e)
                            
                            # Format result
                            result_data = {
                                "transcription": transcription,
                                "response": response.get("response", ""),
                                "speech_audio": speech_audio,
                                "tts_error": tts_error,
                                "confidence": result.confidence,
                                "is_final": True
                            }
                            
                            nonlocal results_count
                            results_count += 1
                            
                            # Call callback if provided
                            if result_callback:
                                await result_callback(result_data)
                
                # Process chunk
                await self.speech_recognizer.process_audio_chunk(chunk, process_result)
                
            # Stop streaming session
            await self.speech_recognizer.stop_streaming()
            
            # Return stats
            return {
                "status": "complete",
                "chunks_processed": chunks_processed,
                "results_count": results_count,
                "total_time": time.time() - start_time
            }
            
        except Exception as e:
            logger.error(f"Error processing streaming audio: {e}")
            await self.speech_recognizer.stop_streaming()
            
            return {
                "status": "error",
                "error": str(e),
                "chunks_processed": chunks_processed,
                "results_count": results_count,
                "total_time": time.time() - start_time
            }
    
    async def add_knowledge_from_text(self, text: str, source: str = "manual_input") -> Dict[str, Any]:
        """
        Add knowledge from text to the knowledge base.
        
        Args:
            text: Text content to add
            source: Source identifier
            
        Returns:
            Result of the operation
        """
        try:
            # Process text into documents
            documents = self.document_processor.process_text(text, source)
            
            # Add to Pinecone
            upserted_count = await self.pinecone_manager.upsert_documents(documents)
            
            logger.info(f"Added {upserted_count} documents from text")
            return {
                "success": True,
                "documents_added": upserted_count,
                "source": source
            }
        except Exception as e:
            logger.error(f"Error adding knowledge from text: {e}")
            return {
                "success": False,
                "error": str(e)
            }
    
    async def add_knowledge_from_file(self, file_path: str) -> Dict[str, Any]:
        """
        Add knowledge from a file to the knowledge base.
        
        Args:
            file_path: Path to the file
            
        Returns:
            Result of the operation
        """
        try:
            # Process file into documents
            documents = self.document_processor.process_file(file_path)
            
            # Add to Pinecone
            upserted_count = await self.pinecone_manager.upsert_documents(documents)
            
            logger.info(f"Added {upserted_count} documents from file {file_path}")
            return {
                "success": True,
                "documents_added": upserted_count,
                "file_path": file_path
            }
        except Exception as e:
            logger.error(f"Error adding knowledge from file: {e}")
            return {
                "success": False,
                "error": str(e),
                "file_path": file_path
            }
    
    async def get_knowledge_stats(self) -> Dict[str, Any]:
        """Get statistics about the knowledge base."""
        try:
            return await self.pinecone_manager.get_stats()
        except Exception as e:
            logger.error(f"Error getting knowledge stats: {e}")
            return {"error": str(e)}
    
    def _process_audio(self, audio: np.ndarray) -> np.ndarray:
        """Process audio for better speech recognition."""
        try:
            from scipy import signal
            
            # Update noise floor from quiet sections
            self._update_noise_floor(audio)
            
            # Apply high-pass filter
            b, a = signal.butter(6, 100/(16000/2), 'highpass')
            audio = signal.filtfilt(b, a, audio)
            
            # Apply band-pass filter for telephony
            b, a = signal.butter(4, [300/(16000/2), 3400/(16000/2)], 'band')
            audio = signal.filtfilt(b, a, audio)
            
            # Apply pre-emphasis
            audio = np.append(audio[0], audio[1:] - 0.97 * audio[:-1])
            
            # Apply noise gate
            threshold = self.noise_floor * 3.0
            audio = np.where(np.abs(audio) < threshold, 0, audio)
            
            # Normalize
            max_val = np.max(np.abs(audio))
            if max_val > 0:
                audio = audio * (0.9 / max_val)
                
            return audio
        except Exception as e:
            logger.error(f"Error processing audio: {e}")
            return audio
    
    def _update_noise_floor(self, audio: np.ndarray) -> None:
        """Update noise floor estimate from quiet sections."""
        frame_size = min(len(audio), int(0.02 * 16000))
        if frame_size <= 1:
            return
            
        frames = [audio[i:i+frame_size] for i in range(0, len(audio), frame_size)]
        frame_energies = [np.mean(np.square(frame)) for frame in frames]
        
        if len(frame_energies) > 0:
            sorted_energies = sorted(frame_energies)
            quiet_count = max(1, len(sorted_energies) // 10)
            quiet_energies = sorted_energies[:quiet_count]
            
            self.noise_samples.extend(quiet_energies)
            
            if len(self.noise_samples) > self.max_noise_samples:
                self.noise_samples = self.noise_samples[-self.max_noise_samples:]
            
            if self.noise_samples:
                self.noise_floor = max(
                    0.001,
                    min(0.02, np.percentile(self.noise_samples, 90) * 1.5)
                )
    
    @property
    def initialized(self) -> bool:
        """Check if all components are initialized."""
        return (self.speech_recognizer is not None and 
                self.conversation_manager is not None and 
                self.query_engine is not None and
                self.tts_client is not None and
                self.pinecone_manager is not None)
    
    async def shutdown(self):
        """Shut down all components properly."""
        logger.info("Shutting down Voice AI Agent...")
        
        # Close Google Cloud streaming session
        if self.speech_recognizer and hasattr(self.speech_recognizer, 'is_streaming') and self.speech_recognizer.is_streaming:
            await self.speech_recognizer.stop_streaming()
        
        # No specific cleanup needed for OpenAI/Pinecone components
        logger.info("Voice AI Agent shutdown complete")
```

## Updated integration/pipeline.py

```python
"""
End-to-end pipeline orchestration - Updated for OpenAI + Pinecone.
"""
import os
import asyncio
import logging
import time
from typing import Optional, Dict, Any, AsyncIterator, Union, List, Callable, Awaitable

import numpy as np

# Updated imports for new knowledge base
from speech_to_text.google_cloud_stt import GoogleCloudStreamingSTT
from speech_to_text.stt_integration import STTIntegration
from knowledge_base.conversation_manager import ConversationManager
from knowledge_base.query_engine import QueryEngine
from integration.tts_integration import TTSIntegration

# Minimum word count for a valid user query
MIN_VALID_WORDS = 2

logger = logging.getLogger(__name__)

class VoiceAIAgentPipeline:
    """
    End-to-end pipeline orchestration - Updated for OpenAI + Pinecone.
    """
    
    def __init__(
        self,
        speech_recognizer: Union[GoogleCloudStreamingSTT, Any],
        conversation_manager: ConversationManager,
        query_engine: QueryEngine,
        tts_integration: TTSIntegration
    ):
        """
        Initialize the pipeline with updated components.
        
        Args:
            speech_recognizer: Initialized STT component  
            conversation_manager: OpenAI conversation manager
            query_engine: OpenAI + Pinecone query engine
            tts_integration: TTS integration with ElevenLabs
        """
        self.speech_recognizer = speech_recognizer
        self.conversation_manager = conversation_manager
        self.query_engine = query_engine
        self.tts_integration = tts_integration
        
        # Create a helper for filtering out non-speech transcriptions
        self.stt_helper = STTIntegration(speech_recognizer)
        
        # Determine if we're using Google Cloud STT
        self.using_google_cloud = isinstance(speech_recognizer, GoogleCloudStreamingSTT)
        logger.info(f"Pipeline initialized with {'Google Cloud' if self.using_google_cloud else 'Other'} STT and OpenAI + Pinecone")
    
    async def _is_valid_transcription(self, transcription: str) -> bool:
        """Check if a transcription is valid and should be processed."""
        cleaned_text = self.stt_helper.cleanup_transcription(transcription)
        
        if not cleaned_text:
            return False
            
        words = cleaned_text.split()
        if len(words) < MIN_VALID_WORDS:
            return False
            
        return True
    
    async def process_audio_file(
        self,
        audio_file_path: str,
        user_id: Optional[str] = None,
        output_speech_file: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Process an audio file through the complete pipeline.
        
        Args:
            audio_file_path: Path to the input audio file
            user_id: User identifier for conversation tracking
            output_speech_file: Path to save the output speech file (optional)
            
        Returns:
            Dictionary with results from each stage
        """
        logger.info(f"Starting pipeline with audio: {audio_file_path}")
        
        # Track timing for each stage
        timings = {}
        start_time = time.time()
        
        # STAGE 1: Speech-to-Text
        logger.info("STAGE 1: Speech-to-Text")
        stt_start = time.time()
        
        # Log audio file info
        logger.info(f"Audio file size: {os.path.getsize(audio_file_path)} bytes")
        
        from speech_to_text.utils.audio_utils import load_audio_file
        
        # Load audio file
        try:
            audio, sample_rate = load_audio_file(audio_file_path, target_sr=16000)
            logger.info(f"Loaded audio: {len(audio)} samples, {sample_rate}Hz")
        except Exception as e:
            logger.error(f"Error loading audio file: {e}", exc_info=True)
            return {"error": f"Error loading audio file: {e}"}
        
        # Process for transcription
        logger.info("Transcribing audio...")
        transcription, duration = await self._transcribe_audio(audio)
        
        # Validate transcription
        is_valid = await self._is_valid_transcription(transcription)
        if not is_valid:
            logger.warning(f"Transcription not valid: '{transcription}'")
            return {"error": "No valid transcription detected", "transcription": transcription}
            
        timings["stt"] = time.time() - stt_start
        logger.info(f"Transcription completed in {timings['stt']:.2f}s: {transcription}")
        
        # STAGE 2: Knowledge Base Query with OpenAI
        logger.info("STAGE 2: Knowledge Base Query with OpenAI + Pinecone")
        kb_start = time.time()
        
        try:
            # Use conversation manager for response generation
            response_result = await self.conversation_manager.handle_user_input(
                user_id=user_id or "default_user",
                message=transcription
            )
            response = response_result.get("response", "")
            
            if not response:
                return {"error": "No response generated from knowledge base"}
                
            timings["kb"] = time.time() - kb_start
            logger.info(f"Response generated: {response[:50]}...")
            
        except Exception as e:
            logger.error(f"Error in KB stage: {e}")
            return {"error": f"Knowledge base error: {str(e)}"}
        
        # STAGE 3: Text-to-Speech with ElevenLabs
        logger.info("STAGE 3: Text-to-Speech with ElevenLabs")
        tts_start = time.time()
        
        try:
            # Convert response to speech
            speech_audio = await self.tts_integration.text_to_speech(response)
            
            # Save speech audio if output file specified
            if output_speech_file:
                os.makedirs(os.path.dirname(os.path.abspath(output_speech_file)), exist_ok=True)
                with open(output_speech_file, "wb") as f:
                    f.write(speech_audio)
                logger.info(f"Saved speech audio to {output_speech_file}")
            
            timings["tts"] = time.time() - tts_start
            logger.info(f"TTS completed in {timings['tts']:.2f}s, generated {len(speech_audio)} bytes")
            
        except Exception as e:
            logger.error(f"Error in TTS stage: {e}")
            return {
                "error": f"TTS error: {str(e)}",
                "transcription": transcription,
                "response": response
            }
        
        # Calculate total time
        total_time = time.time() - start_time
        logger.info(f"Pipeline completed in {total_time:.2f}s")
        
        # Compile results
        return {
            "transcription": transcription,
            "response": response,
            "speech_audio_size": len(speech_audio),
            "speech_audio": None if output_speech_file else speech_audio,
            "timings": timings,
            "total_time": total_time
        }
    
    async def process_audio_streaming(
        self,
        audio_data: Union[bytes, np.ndarray],
        user_id: Optional[str] = None,
        audio_callback: Callable[[bytes], Awaitable[None]] = None
    ) -> Dict[str, Any]:
        """
        Process audio data with streaming response.
        
        Args:
            audio_data: Audio data as bytes or numpy array
            user_id: User identifier
            audio_callback: Callback to handle audio data
            
        Returns:
            Dictionary with stats about the process
        """
        logger.info(f"Starting streaming pipeline with audio: {type(audio_data)}")
        
        # Record start time
        start_time = time.time()
        
        try:
            # Ensure audio is in the right format
            if isinstance(audio_data, bytes):
                audio = np.frombuffer(audio_data, dtype=np.float32)
            else:
                audio = audio_data
            
            # Transcribe audio
            transcription, duration = await self._transcribe_audio(audio)
            
            # Validate transcription
            is_valid = await self._is_valid_transcription(transcription)
            if not is_valid:
                logger.warning(f"Transcription not valid: '{transcription}'")
                return {"error": "No valid transcription detected", "transcription": transcription}
                
            logger.info(f"Transcription: {transcription}")
            transcription_time = time.time() - start_time
        except Exception as e:
            logger.error(f"Error in transcription: {e}")
            return {"error": f"Transcription error: {str(e)}"}
        
        # Stream the response
        try:
            total_chunks = 0
            total_audio_bytes = 0
            response_start_time = time.time()
            full_response = ""
            
            # Use conversation manager's streaming method
            async for chunk in self.conversation_manager.handle_user_input_streaming(
                user_id=user_id or "default_user",
                message=transcription
            ):
                chunk_text = chunk.get("chunk", "")
                
                if chunk_text:
                    # Add to full response
                    full_response += chunk_text
                    
                    # Convert to speech and send to callback
                    audio_data = await self.tts_integration.text_to_speech(chunk_text)
                    if audio_callback:
                        await audio_callback(audio_data)
                    
                    # Update stats
                    total_chunks += 1
                    total_audio_bytes += len(audio_data)
                
                if chunk.get("done", False):
                    if chunk.get("full_response"):
                        full_response = chunk["full_response"]
                    break
            
            # Calculate stats
            response_time = time.time() - response_start_time
            total_time = time.time() - start_time
            
            return {
                "transcription": transcription,
                "transcription_time": transcription_time,
                "response_time": response_time,
                "total_time": total_time,
                "total_chunks": total_chunks,
                "total_audio_bytes": total_audio_bytes,
                "full_response": full_response
            }
            
        except Exception as e:
            logger.error(f"Error in streaming response: {e}")
            return {
                "error": f"Streaming error: {str(e)}",
                "transcription": transcription,
                "transcription_time": transcription_time
            }
    
    async def process_audio_data(
        self,
        audio_data: Union[bytes, np.ndarray],
        user_id: Optional[str] = None,
        speech_output_path: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Process audio data through the complete pipeline.
        
        Args:
            audio_data: Audio data as bytes or numpy array
            user_id: User identifier
            speech_output_path: Path to save speech output
            
        Returns:
            Results dictionary
        """
        logger.info(f"Starting pipeline with audio data: {type(audio_data)}")
        
        # Track timing
        start_time = time.time()
        
        # Convert audio data if needed
        if isinstance(audio_data, bytes):
            audio = np.frombuffer(audio_data, dtype=np.float32)
        else:
            audio = audio_data
        
        # STAGE 1: Speech-to-Text
        logger.info("STAGE 1: Speech-to-Text")
        stt_start = time.time()
        
        # Process for transcription
        transcription, duration = await self._transcribe_audio(audio)
        
        # Validate transcription
        is_valid = await self._is_valid_transcription(transcription)
        if not is_valid:
            logger.warning(f"Transcription not valid: '{transcription}'")
            return {"error": "No valid transcription detected", "transcription": transcription}
            
        timings = {"stt": time.time() - stt_start}
        logger.info(f"Transcription completed in {timings['stt']:.2f}s: {transcription}")
        
        # STAGE 2: Knowledge Base Query
        logger.info("STAGE 2: Knowledge Base Query with OpenAI + Pinecone")
        kb_start = time.time()
        
        try:
            # Use conversation manager for response generation
            response_result = await self.conversation_manager.handle_user_input(
                user_id=user_id or "default_user",
                message=transcription
            )
            response = response_result.get("response", "")
            
            if not response:
                return {"error": "No response generated from knowledge base"}
                
            timings["kb"] = time.time() - kb_start
            logger.info(f"Response generated: {response[:50]}...")
            
        except Exception as e:
            logger.error(f"Error in KB stage: {e}")
            return {"error": f"Knowledge base error: {str(e)}"}
        
        # STAGE 3: Text-to-Speech
        logger.info("STAGE 3: Text-to-Speech with ElevenLabs")
        tts_start = time.time()
        
        try:
            # Convert response to speech
            speech_audio = await self.tts_integration.text_to_speech(response)
            
            # Save speech audio if output file specified
            if speech_output_path:
                os.makedirs(os.path.dirname(os.path.abspath(speech_output_path)), exist_ok=True)
                with open(speech_output_path, "wb") as f:
                    f.write(speech_audio)
                logger.info(f"Saved speech audio to {speech_output_path}")
            
            timings["tts"] = time.time() - tts_start
            logger.info(f"TTS completed in {timings['tts']:.2f}s")
            
            # Calculate total time
            total_time = time.time() - start_time
            
            # Compile results
            return {
                "transcription": transcription,
                "response": response,
                "speech_audio_size": len(speech_audio),
                "speech_audio": speech_audio,
                "timings": timings,
                "total_time": total_time
            }
            
        except Exception as e:
            logger.error(f"Error in TTS stage: {e}")
            return {
                "error": f"TTS error: {str(e)}",
                "transcription": transcription,
                "response": response
            }
    
    async def _transcribe_audio(self, audio: np.ndarray) -> tuple[str, float]:
        """Transcribe audio data using Google Cloud STT."""
        logger.info(f"Transcribing audio: {len(audio)} samples")
        
        if self.using_google_cloud:
            return await self._transcribe_audio_google_cloud(audio)
        else:
            return await self._transcribe_audio_generic(audio)
    
    async def _transcribe_audio_google_cloud(self, audio: np.ndarray) -> tuple[str, float]:
        """Transcribe audio using Google Cloud STT."""
        try:
            # Convert to 16-bit PCM bytes
            audio_bytes = (audio * 32767).astype(np.int16).tobytes()
            
            # Start a streaming session
            await self.speech_recognizer.start_streaming()
            
            # Track final results
            final_results = []
            
            # Process callback to collect results
            async def collect_result(result):
                if result.is_final:
                    final_results.append(result)
            
            # Process audio in chunks
            chunk_size = 4096  # ~128ms at 16kHz
            for i in range(0, len(audio_bytes), chunk_size):
                chunk = audio_bytes[i:i+chunk_size]
                result = await self.speech_recognizer.process_audio_chunk(chunk, collect_result)
                
                # Add final results directly
                if result and result.is_final:
                    final_results.append(result)
            
            # Stop streaming
            transcription, duration = await self.speech_recognizer.stop_streaming()
            
            # If we didn't get a transcription from stop_streaming but have final results
            if not transcription and final_results:
                # Get best final result based on confidence
                best_result = max(final_results, key=lambda r: r.confidence)
                transcription = best_result.text
                duration = best_result.end_time - best_result.start_time if best_result.end_time > 0 else len(audio) / 16000
            
            # Clean up the transcription
            transcription = self.stt_helper.cleanup_transcription(transcription)
            
            return transcription, duration
            
        except Exception as e:
            logger.error(f"Error in Google Cloud transcription: {e}", exc_info=True)
            return "", len(audio) / 16000
    
    async def _transcribe_audio_generic(self, audio: np.ndarray) -> tuple[str, float]:
        """Generic transcription method for any STT system."""
        try:
            # Start streaming
            if hasattr(self.speech_recognizer, 'start_streaming'):
                await self.speech_recognizer.start_streaming()
            
            # Process audio chunk
            if hasattr(self.speech_recognizer, 'process_audio_chunk'):
                await self.speech_recognizer.process_audio_chunk(audio)
            
            # Get final transcription
            if hasattr(self.speech_recognizer, 'stop_streaming'):
                transcription, duration = await self.speech_recognizer.stop_streaming()
            else:
                transcription = ""
                duration = len(audio) / 16000
            
            # Clean up transcription
            if hasattr(self.stt_helper, 'cleanup_transcription'):
                transcription = self.stt_helper.cleanup_transcription(transcription)
            
            return transcription, duration
            
        except Exception as e:
            logger.error(f"Error in generic transcription: {e}", exc_info=True)
            return "", len(audio) / 16000
    
    async def process_realtime_stream(
        self,
        audio_chunk_generator: AsyncIterator[np.ndarray],
        user_id: Optional[str] = None,
        audio_output_callback: Callable[[bytes], Awaitable[None]] = None
    ) -> AsyncIterator[Dict[str, Any]]:
        """
        Process a real-time audio stream.
        
        Args:
            audio_chunk_generator: Async generator producing audio chunks
            user_id: User identifier
            audio_output_callback: Callback to handle output audio data
            
        Yields:
            Status updates and results
        """
        logger.info("Starting real-time audio stream processing")
        
        # Track state
        is_speaking = False
        processing = False
        last_transcription = ""
        silence_frames = 0
        max_silence_frames = 5
        
        # Create audio buffer for processing
        audio_buffer = []
        
        # Timing stats
        start_time = time.time()
        
        try:
            # Initialize the speech recognizer
            if hasattr(self.speech_recognizer, 'start_streaming'):
                await self.speech_recognizer.start_streaming()
            
            # Track results
            results = []
            
            # Define result callback
            async def result_callback(result):
                results.append(result)
                logger.debug(f"Received transcription result: {result.text if hasattr(result, 'text') else str(result)}")
            
            # Process incoming audio chunks
            async for audio_chunk in audio_chunk_generator:
                # Convert if needed
                if isinstance(audio_chunk, bytes):
                    audio_chunk = np.frombuffer(audio_chunk, dtype=np.float32)
                
                # Check for silence
                is_speech = np.mean(np.abs(audio_chunk)) > 0.01  # Simple energy-based detector
                
                if not is_speech:
                    silence_frames += 1
                else:
                    silence_frames = 0
                    
                # Add to buffer
                audio_buffer.append(audio_chunk)
                
                # Process the audio chunk
                if hasattr(self.speech_recognizer, 'process_audio_chunk'):
                    result = await self.speech_recognizer.process_audio_chunk(
                        audio_chunk=audio_chunk,
                        callback=result_callback
                    )
                    
                    # Check for final result
                    if result and hasattr(result, 'is_final') and result.is_final:
                        # Clean up transcription
                        transcription = self.stt_helper.cleanup_transcription(result.text)
                        
                        # Validate transcription
                        if transcription and await self._is_valid_transcription(transcription) and transcription != last_transcription:
                            # Yield status update
                            yield {
                                "status": "transcribed",
                                "transcription": transcription
                            }
                            
                            # Generate response
                            if not is_speaking and not processing:
                                processing = True
                                try:
                                    # Use conversation manager
                                    response_result = await self.conversation_manager.handle_user_input(
                                        user_id=user_id or "default_user",
                                        message=transcription
                                    )
                                    response = response_result.get("response", "")
                                    
                                    if response:
                                        # Mark agent as speaking
                                        is_speaking = True
                                        
                                        # Convert to speech
                                        speech_audio = await self.tts_integration.text_to_speech(response)
                                        
                                        # Send through callback
                                        if audio_output_callback:
                                            await audio_output_callback(speech_audio)
                                        
                                        # Agent is done speaking
                                        is_speaking = False
                                        
                                        # Yield response
                                        yield {
                                            "status": "response",
                                            "transcription": transcription,
                                            "response": response,
                                            "audio_size": len(speech_audio) if speech_audio else 0
                                        }
                                        
                                        # Update last transcription
                                        last_transcription = transcription
                                finally:
                                    processing = False
            
            # Process any final audio
            if hasattr(self.speech_recognizer, 'stop_streaming'):
                final_transcription, _ = await self.speech_recognizer.stop_streaming()
                final_transcription = self.stt_helper.cleanup_transcription(final_transcription)
                
                if final_transcription and await self._is_valid_transcription(final_transcription) and final_transcription != last_transcription:
                    # Generate final response
                    response_result = await self.conversation_manager.handle_user_input(
                        user_id=user_id or "default_user",
                        message=final_transcription
                    )
                    final_response = response_result.get("response", "")
                    
                    if final_response:
                        # Mark agent as speaking
                        is_speaking = True
                        
                        # Convert to speech
                        final_speech = await self.tts_integration.text_to_speech(final_response)
                        
                        # Send through callback
                        if audio_output_callback:
                            await audio_output_callback(final_speech)
                        
                        # Agent is done speaking
                        is_speaking = False
                        
                        # Yield final response
                        yield {
                            "status": "final",
                            "transcription": final_transcription,
                            "response": final_response,
                            "audio_size": len(final_speech) if final_speech else 0,
                            "total_time": time.time() - start_time
                        }
            
            # Yield completion
            yield {
                "status": "complete",
                "total_time": time.time() - start_time
            }
            
        except Exception as e:
            logger.error(f"Error in real-time stream processing: {e}", exc_info=True)
            yield {
                "status": "error",
                "error": str(e),
                "total_time": time.time() - start_time
            }
```

## Migration Script

Create a new file `migrate_to_openai.py` in your project root:

```python
"""
Migration script to convert from LlamaIndex/Ollama to OpenAI + Pinecone.
"""
import asyncio
import logging
import os
from datetime import datetime

# Import current system components
import chromadb
from llama_index.core.schema import Document as LlamaDocument

# Import new components
from knowledge_base.pinecone_manager import PineconeManager
from knowledge_base.document_processor import DocumentProcessor
from knowledge_base.openai_assistant_manager import OpenAIAssistantManager

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class KnowledgeBaseMigration:
    """Migrate knowledge base from LlamaIndex to OpenAI + Pinecone."""
    
    def __init__(self):
        self.pinecone_manager = PineconeManager()
        self.doc_processor = DocumentProcessor()
        self.openai_manager = OpenAIAssistantManager()
    
    async def migrate_from_chromadb(self, chroma_collection_name: str = "company_knowledge"):
        """Migrate documents from ChromaDB to Pinecone."""
        logger.info("Starting migration from ChromaDB to Pinecone...")
        
        # Connect to ChromaDB
        chroma_client = chromadb.Client()
        collection = chroma_client.get_collection(chroma_collection_name)
        
        # Get all documents
        results = collection.get(
            include=["metadatas", "documents", "embeddings"]
        )
        
        logger.info(f"Found {len(results['ids'])} documents in ChromaDB")
        
        # Convert to new format
        documents = []
        for i, (doc_id, metadata, document) in enumerate(zip(
            results['ids'], results['metadatas'], results['documents']
        )):
            # Extract text content
            text = document if isinstance(document, str) else metadata.get('text', '')
            
            # Create document object
            doc = {
                "id": doc_id,
                "text": text,
                "source": metadata.get("source", "unknown"),
                "document_id": metadata.get("document_id", doc_id),
                "metadata": metadata
            }
            documents.append(doc)
        
        # Initialize Pinecone and upload
        await self.pinecone_manager.init()
        upserted_count = await self.pinecone_manager.upsert_documents(documents)
        
        logger.info(f"Migrated {upserted_count} documents to Pinecone")
    
    async def migrate_documents_directory(self, directory_path: str):
        """Migrate documents from a directory to the new system."""
        logger.info(f"Migrating documents from {directory_path}...")
        
        await self.pinecone_manager.init()
        
        # Process each file in directory
        import os
        processed_count = 0
        
        for filename in os.listdir(directory_path):
            file_path = os.path.join(directory_path, filename)
            
            if os.path.isfile(file_path):
                try:
                    # Process file
                    documents = self.doc_processor.process_file(file_path)
                    
                    # Upload to Pinecone
                    await self.pinecone_manager.upsert_documents(documents)
                    
                    processed_count += len(documents)
                    logger.info(f"Processed {filename}: {len(documents)} chunks")
                    
                except Exception as e:
                    logger.error(f"Error processing {filename}: {e}")
        
        logger.info(f"Migration complete. Processed {processed_count} document chunks")
    
    async def create_assistant(self):
        """Create the OpenAI assistant."""
        logger.info("Creating OpenAI Assistant...")
        
        assistant_id = await self.openai_manager.create_assistant()
        logger.info(f"Created assistant: {assistant_id}")
        
        return assistant_id
    
    async def run_migration(self, 
                          migrate_chromadb: bool = True,
                          chroma_collection: str = "company_knowledge",
                          documents_directory: Optional[str] = None):
        """Run the complete migration process."""
        logger.info("Starting complete migration process...")
        
        # Step 1: Migrate from ChromaDB if requested
        if migrate_chromadb:
            try:
                await self.migrate_from_chromadb(chroma_collection)
            except Exception as e:
                logger.error(f"Error migrating from ChromaDB: {e}")
        
        # Step 2: Migrate documents from directory if provided
        if documents_directory and os.path.exists(documents_directory):
            try:
                await self.migrate_documents_directory(documents_directory)
            except Exception as e:
                logger.error(f"Error migrating documents: {e}")
        
        # Step 3: Create OpenAI Assistant
        try:
            assistant_id = await self.create_assistant()
            logger.info(f"Migration complete! Assistant ID: {assistant_id}")
        except Exception as e:
            logger.error(f"Error creating assistant: {e}")
        
        # Step 4: Verify migration
        await self.verify_migration()
    
    async def verify_migration(self):
        """Verify the migration was successful."""
        logger.info("Verifying migration...")
        
        try:
            # Check Pinecone stats
            stats = await self.pinecone_manager.get_stats()
            logger.info(f"Pinecone stats: {stats}")
            
            # Test a query
            results = await self.pinecone_manager.query("test query", top_k=1)
            logger.info(f"Test query returned {len(results)} results")
            
            logger.info("Migration verification complete")
            
        except Exception as e:
            logger.error(f"Error in verification: {e}")

async def main():
    """Main migration function."""
    print("Knowledge Base Migration - ChromaDB/LlamaIndex to OpenAI + Pinecone")
    print("=" * 60)
    
    # Create migration instance
    migration = KnowledgeBaseMigration()
    
    # Run migration
    await migration.run_migration(
        migrate_chromadb=True,
        chroma_collection="company_knowledge",
        documents_directory="./knowledge_base/knowledge_docs"  # Adjust path as needed
    )
    
    print("\nMigration complete!")
    print("Please update your application to use the new knowledge base components.")

if __name__ == "__main__":
    asyncio.run(main())
```

## Environment Variables

Update your `.env` file to include the new required variables:

```env
# OpenAI Configuration
OPENAI_API_KEY=your_openai_api_key_here
OPENAI_MODEL=gpt-4o
OPENAI_EMBEDDING_MODEL=text-embedding-3-small
OPENAI_MAX_TOKENS=4096
OPENAI_TEMPERATURE=0.7

# Pinecone Configuration
PINECONE_API_KEY=your_pinecone_api_key_here
PINECONE_INDEX_NAME=voice-ai-knowledge
PINECONE_ENVIRONMENT=us-east-1-aws

# Redis Configuration (for caching)
REDIS_URL=redis://localhost:6379/0

# Existing configurations remain the same
GOOGLE_APPLICATION_CREDENTIALS=path/to/google/credentials.json
ELEVENLABS_API_KEY=your_elevenlabs_api_key
TTS_VOICE_ID=your_voice_id
# ... other existing env vars
```

## Summary

This implementation provides:

1. **Complete replacement** of the knowledge base folder with OpenAI + Pinecone components
2. **Seamless integration** with your existing speech-to-text and text-to-speech systems
3. **Streaming support** for real-time voice conversations
4. **Caching and rate limiting** for production use
5. **Easy migration script** to transfer existing data
6. **Backward-compatible interfaces** so your existing pipeline code needs minimal changes

To implement:

1. Replace the `knowledge_base/` folder with the new implementation
2. Update `requirements.txt` with the new dependencies
3. Set up your OpenAI and Pinecone API keys in `.env`
4. Run the migration script to transfer existing data
5. Update your main application to use the new components

The new system will maintain the same voice conversation capabilities while leveraging OpenAI's more powerful models and Pinecone's scalable vector storage.